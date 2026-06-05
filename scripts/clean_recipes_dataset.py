"""Clean Kaggle recipe datasets into MongoDB-ready JSON.

Usage:
  python scripts/clean_recipes_dataset.py --input recipes.json --output cleaned_recipes.json
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path
from typing import Any

sys.path.append(str(Path(__file__).resolve().parents[1]))

from app.utils.recipe_utils import clean_text, normalize_text, recipe_document, split_listish


NAME_KEYS = ("tarif_adi", "name", "title", "recipe_name", "yemek_adi")
CATEGORY_KEYS = ("kategori", "category", "course", "type")
INGREDIENT_KEYS = ("malzemeler", "ingredients", "ingredient_list")
STEP_KEYS = ("yapilis_adimlari", "steps", "directions", "instructions", "method", "preparation")
SERVING_KEYS = ("porsiyon", "servings", "portion", "yield")
IMAGE_KEYS = ("image", "image_url", "gorsel", "resim")


def pick(row: dict[str, Any], keys: tuple[str, ...], default: Any = "") -> Any:
    lower_map = {str(key).lower().strip(): key for key in row.keys()}
    for key in keys:
        real_key = lower_map.get(key)
        if real_key is not None and row.get(real_key) not in (None, ""):
            return row.get(real_key)
    return default


def read_dataset(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(data, list):
            return [item for item in data if isinstance(item, dict)]
        if isinstance(data, dict):
            for key in ("items", "recipes", "data"):
                if isinstance(data.get(key), list):
                    return [item for item in data[key] if isinstance(item, dict)]
            return [data]
        return []
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    if suffix in {".xlsx", ".xls"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX okumak icin openpyxl gerekli: pip install openpyxl") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [clean_text(value) for value in rows[0]]
        return [
            {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
            for row in rows[1:]
        ]
    raise ValueError(f"Desteklenmeyen dataset formati: {suffix}")


def duration_from_row(row: dict[str, Any]) -> str:
    duration = pick(row, ("duration", "sure", "total_time", "toplam_sure"), "")
    if duration:
        return clean_text(duration)
    prep = pick(row, ("hazirlik_suresi_dk", "prep_time", "preparation_time"), "")
    cook = pick(row, ("pisirme_suresi_dk", "cook_time", "cooking_time"), "")
    parts = []
    if prep not in (None, ""):
        parts.append(f"Hazirlik {prep} dk")
    if cook not in (None, ""):
        parts.append(f"Pisirme {cook} dk")
    return ", ".join(parts)


def clean_dataset(input_path: Path) -> list[dict[str, Any]]:
    rows = read_dataset(input_path)
    cleaned: list[dict[str, Any]] = []
    seen: set[str] = set()

    for row in rows:
        name = pick(row, NAME_KEYS)
        normalized_name = normalize_text(name)
        if not normalized_name or normalized_name in seen:
            continue

        ingredients_value = pick(row, INGREDIENT_KEYS, [])
        steps_value = pick(row, STEP_KEYS, [])
        ingredients = ingredients_value if isinstance(ingredients_value, list) else split_listish(ingredients_value)
        steps = steps_value if isinstance(steps_value, list) else split_listish(steps_value)
        if not ingredients or not steps:
            continue

        doc = recipe_document(
            name=name,
            category=pick(row, CATEGORY_KEYS, ""),
            ingredients=ingredients,
            steps=steps,
            duration=duration_from_row(row),
            servings=pick(row, SERVING_KEYS, ""),
            image=pick(row, IMAGE_KEYS, ""),
            source="kaggle",
        )
        if not doc["ingredients"] or not doc["steps"]:
            continue
        seen.add(normalized_name)
        cleaned.append(doc)

    return cleaned


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="CSV/JSON/XLSX dataset path")
    parser.add_argument("--output", default="cleaned_recipes.json", help="Output JSON path")
    args = parser.parse_args()

    cleaned = clean_dataset(Path(args.input))
    Path(args.output).write_text(json.dumps(cleaned, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Cleaned {len(cleaned)} recipes -> {args.output}")


if __name__ == "__main__":
    main()
